"""Excel export for manual MES upload."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
import pandas as pd

from data.config import MES_EXPORT_COLUMNS


def build_mes_dataframe(rows: list[dict[str, str | int]]) -> pd.DataFrame:
    data = [
        {
            "小分切条码": row["SR Barcode"],
            "涂覆条码": row["Master Roll SN"],
            "涂覆机台": row["Coating Machine Code"],
            "涂覆日期": row["Coating Date Code"],
            "涂覆班组": row["Coating Team Code"],
            "涂覆流水": row["Coating Sequence"],
            "小分切工位": row["Slitting Position"],
            "小分切机台": row["Slitter Machine Code"],
        }
        for row in rows
    ]
    return pd.DataFrame(data, columns=MES_EXPORT_COLUMNS).astype(str)


def export_mes_excel(rows: list[dict[str, str | int]]) -> bytes:
    df = build_mes_dataframe(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "MES Upload"

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)

    for col_index, column_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_index, value=column_name)
        cell.fill = header_fill
        cell.font = header_font

    for row_index, record in enumerate(df.to_dict("records"), start=2):
        for col_index, column_name in enumerate(df.columns, start=1):
            cell = ws.cell(row=row_index, column=col_index, value=str(record[column_name]))
            cell.number_format = "@"

    for col_index, column_name in enumerate(df.columns, start=1):
        max_length = max(
            len(str(column_name)),
            *(len(str(value)) for value in df[column_name].tolist()),
        )
        ws.column_dimensions[get_column_letter(col_index)].width = min(max_length + 4, 40)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
