import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest


class GenerationWorkflowTests(unittest.TestCase):
    def test_generation_requires_click_and_parameter_changes_clear_result(self):
        app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "app.py")).run()
        app.number_input[0].set_value(100.0)
        app.number_input[1].set_value(1)
        app.button[0].click().run()
        self.assertTrue(app.button(key="generate_batch").disabled)

        app.text_area[0].set_value("ULE3U1B260824UA011").run()
        self.assertTrue(app.button(key="generate_batch").disabled)
        app.button(key="apply_master_rolls").click().run()
        self.assertFalse(app.button(key="generate_batch").disabled)
        self.assertNotIn("generated_result", app.session_state)
        self.assertEqual(len(app.get("download_button")), 0)

        app.button(key="generate_batch").click().run()
        self.assertFalse(app.exception)
        first_result = app.session_state["generated_result"]
        self.assertEqual(first_result["rows"][0]["Slitting Sequence"], "00001")
        self.assertEqual(len(app.get("download_button")), 1)
        app.run()
        self.assertEqual(app.session_state["generated_result"]["excel"], first_result["excel"])

        sequence = next(item for item in app.number_input if item.label == "Starting Slitting Sequence")
        sequence.set_value(200).run()
        self.assertNotIn("generated_result", app.session_state)
        self.assertEqual(len(app.get("download_button")), 0)

        app.button(key="generate_batch").click().run()
        self.assertFalse(app.exception)
        result = app.session_state["generated_result"]
        self.assertEqual(result["rows"][0]["Slitting Sequence"], "00200")
        worksheet = load_workbook(BytesIO(result["excel"])).active
        self.assertEqual(
            [cell.value for cell in worksheet[1]],
            ["小分切条码", "涂覆条码", "涂覆机台", "涂覆日期", "涂覆班组", "涂覆流水", "小分切工位", "小分切机台"],
        )
        self.assertEqual(
            [cell.value for cell in worksheet[2]][1:],
            ["ULE3U1B260824UA011", "02", "BV24", "U", "011", "01", "X01"],
        )
        self.assertTrue(worksheet["A2"].value.endswith("-00200"))
        preview = app.dataframe[-1].value
        self.assertEqual(list(preview.columns), [cell.value for cell in worksheet[1]])
        self.assertEqual(
            list(preview.itertuples(index=False, name=None)),
            list(worksheet.iter_rows(min_row=2, values_only=True)),
        )

        app.text_area[0].set_value("ULE3U1B260824UA012").run()
        self.assertTrue(app.button(key="generate_batch").disabled)
        self.assertNotIn("generated_result", app.session_state)
        self.assertEqual(len(app.get("download_button")), 0)
        app.button(key="apply_master_rolls").click().run()
        self.assertFalse(app.button(key="generate_batch").disabled)
        app.button(key="generate_batch").click().run()
        self.assertEqual(app.session_state["generated_result"]["rows"][0]["Master Roll SN"], "ULE3U1B260824UA012")

        app.text_area[0].set_value("invalid").run()
        app.button(key="apply_master_rolls").click().run()
        self.assertTrue(app.button(key="generate_batch").disabled)
        self.assertTrue(app.error)
        self.assertNotIn("applied_master_roll_inputs", app.session_state)
        self.assertNotIn("generated_result", app.session_state)
        self.assertEqual(len(app.get("download_button")), 0)
        self.assertFalse(app.exception)


if __name__ == "__main__":
    unittest.main()
